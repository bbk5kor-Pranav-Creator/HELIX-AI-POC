import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NumberObject,
)

import document_extractor


def _write_docx(path):
    document = Document()
    document.add_heading("Operations", level=1)
    document.add_paragraph("First paragraph: café intake.")
    document.add_heading("1.1 Review", level=2)
    document.add_paragraph("Review order before approval.")
    first_table = document.add_table(rows=2, cols=2)
    first_table.cell(0, 0).text = "Activity"
    first_table.cell(0, 1).text = "Owner"
    first_table.cell(1, 0).text = "Check amount"
    first_table.cell(1, 1).text = "Zoë"
    document.add_paragraph("An exception applies after review.")
    second_table = document.add_table(rows=1, cols=2)
    second_table.cell(0, 0).text = "Condition"
    second_table.cell(0, 1).text = "Amount > €10,000"
    document.save(path)


def _pdf_with_text_and_blank_scanned_page():
    writer = PdfWriter()
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    font_ref = writer._add_object(font)
    page = writer.add_blank_page(width=612, height=792)
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})
    })
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 14 Tf 50 740 Td (1. Approval) Tj "
        b"0 -20 Td (Manager approves the order.) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)

    writer.add_blank_page(width=612, height=792)
    scanned_page = writer.add_blank_page(width=612, height=792)
    image = DecodedStreamObject()
    image.set_data(bytes([255, 255, 255]))
    image.update({
        NameObject("/Type"): NameObject("/XObject"),
        NameObject("/Subtype"): NameObject("/Image"),
        NameObject("/Width"): NumberObject(1),
        NameObject("/Height"): NumberObject(1),
        NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
        NameObject("/BitsPerComponent"): NumberObject(8),
    })
    image_ref = writer._add_object(image)
    scanned_page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): image_ref})
    })
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


class DocumentExtractorTests(unittest.TestCase):
    def test_extracts_utf8_txt_without_changing_content(self):
        text = "Process overview\r\nStep one.\r\n"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "process.txt"
            path.write_bytes(text.encode("utf-8"))
            self.assertEqual(document_extractor.extract_document(path), text)

    def test_extracts_legacy_text_using_windows_encoding(self):
        text = "Café process"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "process.txt"
            path.write_bytes(text.encode("cp1252"))
            self.assertEqual(document_extractor.extract_document(path), text)

    def test_rejects_empty_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "empty.txt"
            path.write_bytes(b" \r\n\t")
            with self.assertRaisesRegex(
                document_extractor.DocumentExtractionError, "No readable text"
            ):
                document_extractor.extract_document(path)

    def test_rejects_binary_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "binary.txt"
            path.write_bytes(b"not readable\x00text")
            with self.assertRaisesRegex(
                document_extractor.DocumentExtractionError, "binary data"
            ):
                document_extractor.extract_document(path)

    def test_reports_unsupported_format(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "process.odt"
            path.write_bytes(b"content")
            with self.assertRaisesRegex(
                document_extractor.DocumentExtractionError, "Unsupported"
            ):
                document_extractor.extract_document(path)

    def test_legacy_errors_name_each_format_when_soffice_is_unavailable(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch.object(document_extractor.shutil, "which", return_value=None):
                for extension in (".doc", ".xls", ".ppt"):
                    path = Path(directory) / f"legacy{extension}"
                    path.write_bytes(b"legacy")
                    with self.subTest(extension=extension):
                        with self.assertRaisesRegex(
                            document_extractor.DocumentExtractionError,
                            f"Legacy {extension[1:].upper()} files require LibreOffice",
                        ):
                            document_extractor.extract_document(path)

    def test_docx_structured_output_preserves_heading_order_tables_and_unicode(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "operations.docx"
            _write_docx(path)
            structured = document_extractor.extract_structured_document(path)
            blocks = structured["blocks"]

            self.assertEqual(structured["format"], "docx")
            self.assertEqual(
                structured["original_text"],
                document_extractor.extract_document(path),
            )
            self.assertEqual(
                [(block["block_type"], block["text"]) for block in blocks],
                [
                    ("heading", "Operations"),
                    ("paragraph", "First paragraph: café intake."),
                    ("heading", "1.1 Review"),
                    ("paragraph", "Review order before approval."),
                    ("table", ""),
                    ("table_row", "Activity | Owner"),
                    ("table_row", "Check amount | Zoë"),
                    ("paragraph", "An exception applies after review."),
                    ("table", ""),
                    ("table_row", "Condition | Amount > €10,000"),
                ],
            )
            self.assertEqual(
                [blocks[0]["heading_level"], blocks[2]["heading_level"]],
                [1, 2],
            )
            self.assertEqual(blocks[5]["parent_block_ordinal"], blocks[4]["ordinal"])
            self.assertEqual(
                [cell["text"] for cell in blocks[6]["cells"]],
                ["Check amount", "Zoë"],
            )
            self.assertEqual(blocks[1]["paragraph_index"], 1)
            self.assertIsNone(blocks[1]["page_start"])
            self.assertEqual(
                [block["table_index"] for block in blocks if block["block_type"] == "table"],
                [0, 1],
            )

    def test_docx_without_heading_styles_is_preserved_and_warned(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "plain.docx"
            document = Document()
            document.add_paragraph("First unheaded paragraph.")
            document.add_paragraph("Second unheaded paragraph.")
            document.save(path)

            result = document_extractor.extract_structured_document(path)
            self.assertEqual(
                [block["block_type"] for block in result["blocks"]],
                ["paragraph", "paragraph"],
            )
            self.assertIn("no_headings_detected", {
                warning["code"] for warning in result["warnings"]
            })

    def test_txt_preserves_paragraph_boundaries_reading_order_and_heading_levels(self):
        text = "# Process Overview\nFirst paragraph: café.\n\n## Approval\nSecond paragraph: Zoë ✓.\n"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "process.txt"
            path.write_bytes(text.encode("utf-8"))
            result = document_extractor.extract_structured_document(path)

            self.assertEqual(result["original_text"], text)
            self.assertEqual(
                [(block["block_type"], block["text"]) for block in result["blocks"]],
                [
                    ("heading", "# Process Overview"),
                    ("paragraph", "First paragraph: café."),
                    ("heading", "## Approval"),
                    ("paragraph", "Second paragraph: Zoë ✓."),
                ],
            )
            self.assertEqual(
                [result["blocks"][0]["heading_level"], result["blocks"][2]["heading_level"]],
                [1, 2],
            )
            self.assertEqual(result["blocks"][3]["locator"]["line_start"], 5)

    def test_txt_without_headings_remains_ordered_and_warned(self):
        text = "First paragraph.\n\nSecond paragraph with café."
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "unheaded.txt"
            path.write_bytes(text.encode("utf-8"))
            result = document_extractor.extract_structured_document(path)

            self.assertEqual(result["original_text"], text)
            self.assertEqual(
                [block["text"] for block in result["blocks"]],
                ["First paragraph.", "Second paragraph with café."],
            )
            self.assertIn("no_headings_detected", {
                warning["code"] for warning in result["warnings"]
            })

    def test_pdf_structured_output_keeps_page_provenance_and_flags_scanned_page(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "approval.pdf"
            path.write_bytes(_pdf_with_text_and_blank_scanned_page())
            result = document_extractor.extract_structured_document(path)

            self.assertIn("--- Page 1 ---", result["original_text"])
            self.assertIn("--- Page 2 ---", result["original_text"])
            self.assertIn("--- Page 3 ---", result["original_text"])
            self.assertEqual(
                [block["page_start"] for block in result["blocks"]],
                [1, 1],
            )
            self.assertEqual(
                [block["text"] for block in result["blocks"]],
                ["1. Approval", "Manager approves the order."],
            )
            warning_by_code = {
                (warning["code"], warning.get("page_number"))
                for warning in result["warnings"]
            }
            self.assertIn(("page_no_extractable_text", 2), warning_by_code)
            self.assertIn(("page_no_extractable_text", 3), warning_by_code)
            self.assertIn(("suspected_scanned_page", 3), warning_by_code)

    def test_legacy_doc_structured_conversion_preserves_structure_and_warns(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "legacy.doc"
            path.write_bytes(b"legacy DOC input")

            def fake_conversion(command, **_kwargs):
                output_directory = Path(command[command.index("--outdir") + 1])
                converted_path = output_directory / "legacy.docx"
                document = Document()
                document.add_heading("Legacy heading", level=1)
                document.add_paragraph("Legacy paragraph.")
                document.save(converted_path)
                return type("Completed", (), {
                    "returncode": 0,
                    "stdout": "",
                    "stderr": "",
                })()

            with patch.object(
                document_extractor,
                "_soffice_executable",
                return_value="soffice",
            ), patch.object(
                document_extractor.subprocess,
                "run",
                side_effect=fake_conversion,
            ) as conversion:
                result = document_extractor.extract_structured_document(path)

            self.assertEqual(result["format"], "doc")
            self.assertEqual(
                [block["block_type"] for block in result["blocks"]],
                ["heading", "paragraph"],
            )
            self.assertTrue(all(block["page_start"] is None for block in result["blocks"]))
            self.assertIn("legacy_format_converted", {
                warning["code"] for warning in result["warnings"]
            })
            self.assertEqual(conversion.call_args.kwargs["timeout"],
                             document_extractor.CONVERSION_TIMEOUT_SECONDS)

    def test_structured_doc_reports_clear_error_without_libreoffice(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "legacy.doc"
            path.write_bytes(b"legacy DOC input")
            with patch.object(document_extractor, "_soffice_executable", return_value=None):
                with self.assertRaisesRegex(
                    document_extractor.DocumentExtractionError,
                    "Legacy DOC files require LibreOffice",
                ):
                    document_extractor.extract_structured_document(path)

    def test_structured_api_rejects_non_docupedia_formats(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "spreadsheet.xlsx"
            path.write_bytes(b"not used")
            with self.assertRaisesRegex(
                document_extractor.DocumentExtractionError,
                "Unsupported Docupedia document format",
            ):
                document_extractor.extract_structured_document(path)


if __name__ == "__main__":
    unittest.main()
