import sqlite3
import tempfile
import unittest
from pathlib import Path

from docx import Document

import document_extractor
import knowledge_store


class DocupediaPersistenceTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.database = Path(self.directory.name) / "knowledge.db"
        self.original_database = knowledge_store.DATABASE_PATH
        knowledge_store.DATABASE_PATH = self.database
        knowledge_store.initialize_database()
        self.addCleanup(self._restore_database)

    def tearDown(self):
        self.directory.cleanup()

    def _restore_database(self):
        knowledge_store.DATABASE_PATH = self.original_database

    def _write_docx(self, path):
        document = Document()
        document.add_heading("Operations", level=1)
        document.add_paragraph("The intake team records each request.")
        document.add_heading("Review", level=2)
        document.add_paragraph("A manager approves requests over €500.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Activity"
        table.cell(0, 1).text = "Owner"
        table.cell(1, 0).text = "Validate request"
        table.cell(1, 1).text = "Manager"
        document.save(path)

    def test_round_trip_preserves_extraction_order_structure_and_provenance(self):
        path = Path(self.directory.name) / "operations.docx"
        self._write_docx(path)
        extracted = document_extractor.extract_structured_document(path)

        source = knowledge_store.save_structured_document(
            extracted, filename="operations.docx"
        )
        reconstructed = knowledge_store.get_structured_document(source["id"])

        self.assertEqual(reconstructed, extracted)
        self.assertEqual(
            [block["ordinal"] for block in reconstructed["blocks"]],
            list(range(len(reconstructed["blocks"]))),
        )
        self.assertEqual(
            [block["block_type"] for block in reconstructed["blocks"]],
            [
                "heading", "paragraph", "heading", "paragraph", "table",
                "table_row", "table_row",
            ],
        )
        row = reconstructed["blocks"][5]
        self.assertEqual(row["parent_block_ordinal"], 4)
        self.assertEqual(
            [cell["text"] for cell in row["cells"]],
            ["Activity", "Owner"],
        )
        self.assertIsNone(reconstructed["blocks"][1]["page_start"])
        self.assertEqual(
            [block["heading_level"] for block in reconstructed["blocks"]
             if block["block_type"] == "heading"],
            [1, 2],
        )
        with knowledge_store._connection() as connection:
            stored_blocks = connection.execute(
                """
                SELECT ordinal, parent_block_id, parent_heading_block_id
                FROM document_blocks WHERE source_id = ? ORDER BY ordinal
                """,
                (source["id"],),
            ).fetchall()
            block_ids = {
                row["ordinal"]: connection.execute(
                    "SELECT id FROM document_blocks WHERE source_id = ? AND ordinal = ?",
                    (source["id"], row["ordinal"]),
                ).fetchone()["id"]
                for row in stored_blocks
            }
            self.assertEqual(stored_blocks[2]["parent_heading_block_id"], block_ids[0])
            self.assertEqual(stored_blocks[3]["parent_heading_block_id"], block_ids[2])
            self.assertEqual(stored_blocks[5]["parent_block_id"], block_ids[4])
            section_parents = connection.execute(
                """
                SELECT title, parent_section_id FROM document_sections
                WHERE source_id = ? ORDER BY ordinal
                """,
                (source["id"],),
            ).fetchall()
            self.assertIsNone(section_parents[0]["parent_section_id"])
            self.assertIsNotNone(section_parents[1]["parent_section_id"])

            source_row = connection.execute(
                "SELECT original_text, source_type FROM knowledge_sources WHERE id = ?",
                (source["id"],),
            ).fetchone()
            raw_record = connection.execute(
                "SELECT original_text FROM document_records WHERE source_id = ?",
                (source["id"],),
            ).fetchone()
            self.assertEqual(source_row["original_text"], "")
            self.assertEqual(source_row["source_type"], "Docupedia")
            self.assertEqual(raw_record["original_text"], extracted["original_text"])
            self.assertNotEqual(source_row["original_text"], raw_record["original_text"])

    def test_sections_and_chunks_are_heading_aware_and_link_source_blocks(self):
        body = "Intake review decision " * 180
        structured = {
            "format": "txt",
            "original_text": f"# Intake\n{body}\n# Exceptions\nEscalate urgent requests.",
            "warnings": [],
            "blocks": [
                {
                    "ordinal": 0,
                    "block_type": "heading",
                    "text": "# Intake",
                    "heading_level": 1,
                    "page_start": None,
                    "page_end": None,
                    "paragraph_index": 0,
                    "locator": {"line_start": 1, "line_end": 1},
                },
                {
                    "ordinal": 1,
                    "block_type": "paragraph",
                    "text": body,
                    "heading_level": None,
                    "page_start": None,
                    "page_end": None,
                    "paragraph_index": 1,
                    "locator": {"line_start": 2, "line_end": 2},
                },
                {
                    "ordinal": 2,
                    "block_type": "heading",
                    "text": "# Exceptions",
                    "heading_level": 1,
                    "page_start": None,
                    "page_end": None,
                    "paragraph_index": 2,
                    "locator": {"line_start": 3, "line_end": 3},
                },
                {
                    "ordinal": 3,
                    "block_type": "paragraph",
                    "text": "Escalate urgent requests.",
                    "heading_level": None,
                    "page_start": None,
                    "page_end": None,
                    "paragraph_index": 3,
                    "locator": {"line_start": 4, "line_end": 4},
                },
            ],
        }

        source = knowledge_store.save_structured_document(
            structured, filename="workflow.txt"
        )
        with knowledge_store._connection() as connection:
            sections = connection.execute(
                """
                SELECT id, title, parent_section_id, first_block_ordinal,
                       last_block_ordinal, is_synthetic
                FROM document_sections WHERE source_id = ? ORDER BY ordinal
                """,
                (source["id"],),
            ).fetchall()
            self.assertEqual([row["title"] for row in sections], ["# Intake", "# Exceptions"])
            self.assertEqual([row["parent_section_id"] for row in sections], [None, None])
            self.assertEqual(
                [(row["first_block_ordinal"], row["last_block_ordinal"])
                 for row in sections],
                [(0, 1), (2, 3)],
            )
            self.assertTrue(all(row["is_synthetic"] == 0 for row in sections))

            chunks = connection.execute(
                """
                SELECT c.id, c.content, s.title AS section_title
                FROM knowledge_chunks AS c
                JOIN knowledge_chunk_sections AS cs ON cs.chunk_id = c.id
                JOIN document_sections AS s ON s.id = cs.section_id
                WHERE c.source_id = ? ORDER BY c.chunk_index, cs.ordinal
                """,
                (source["id"],),
            ).fetchall()
            self.assertGreater(len(chunks), 2)
            for chunk in chunks:
                self.assertIn(chunk["section_title"], {"# Intake", "# Exceptions"})
                expected_section_content = (
                    "Intake" if chunk["section_title"] == "# Intake"
                    else "Escalate urgent requests."
                )
                self.assertIn(expected_section_content, chunk["content"])

            linked = connection.execute(
                """
                SELECT c.chunk_index, b.ordinal, b.text, cb.start_offset, cb.end_offset
                FROM knowledge_chunk_blocks AS cb
                JOIN knowledge_chunks AS c ON c.id = cb.chunk_id
                JOIN document_blocks AS b ON b.id = cb.block_id
                WHERE cb.source_id = ? ORDER BY c.chunk_index, cb.ordinal
                """,
                (source["id"],),
            ).fetchall()
            self.assertTrue(linked)
            self.assertTrue(all(row["end_offset"] >= row["start_offset"] for row in linked))
            self.assertTrue(any(row["ordinal"] == 1 for row in linked))
            self.assertTrue(any(row["ordinal"] == 3 for row in linked))

    def test_unheaded_document_gets_synthetic_fallback_section_and_chunk(self):
        structured = {
            "format": "txt",
            "original_text": "The unclassified process text remains searchable.",
            "warnings": [{"code": "no_headings_detected"}],
            "blocks": [
                {
                    "ordinal": 0,
                    "block_type": "paragraph",
                    "text": "The unclassified process text remains searchable.",
                    "page_start": None,
                    "page_end": None,
                    "paragraph_index": 0,
                    "locator": {"line_start": 1, "line_end": 1},
                }
            ],
        }
        source = knowledge_store.save_structured_document(
            structured, filename="unheaded.txt"
        )
        with knowledge_store._connection() as connection:
            section = connection.execute(
                "SELECT title, is_synthetic, first_block_ordinal, last_block_ordinal "
                "FROM document_sections WHERE source_id = ?",
                (source["id"],),
            ).fetchone()
            chunk = connection.execute(
                "SELECT content FROM knowledge_chunks WHERE source_id = ?",
                (source["id"],),
            ).fetchone()
            self.assertEqual(section["title"], "Document content")
            self.assertEqual(section["is_synthetic"], 1)
            self.assertEqual((section["first_block_ordinal"], section["last_block_ordinal"]),
                             (0, 0))
            self.assertEqual(
                chunk["content"],
                "The unclassified process text remains searchable.",
            )

    def test_failed_ingestion_rolls_back_source_and_all_document_rows(self):
        structured = {
            "format": "txt",
            "original_text": "Rollback this source content.",
            "blocks": [{
                "ordinal": 0,
                "block_type": "paragraph",
                "text": "Rollback this source content.",
            }],
            "warnings": [],
        }
        with knowledge_store._connection() as connection:
            connection.execute(
                """
                CREATE TRIGGER reject_docupedia_block BEFORE INSERT ON document_blocks
                BEGIN
                    SELECT RAISE(ABORT, 'test ingestion failure');
                END
                """
            )

        with self.assertRaisesRegex(sqlite3.IntegrityError, "test ingestion failure"):
            knowledge_store.save_structured_document(
                structured, filename="rollback.txt"
            )

        with knowledge_store._connection() as connection:
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) FROM knowledge_sources WHERE source_type = 'Docupedia'"
                ).fetchone()[0],
                0,
            )
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM document_records").fetchone()[0],
                0,
            )
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM knowledge_chunks").fetchone()[0],
                0,
            )

    def test_pending_docupedia_chunks_are_not_employee_search_results(self):
        structured = {
            "format": "txt",
            "original_text": "Pendingquartz policy wording.",
            "blocks": [{
                "ordinal": 0,
                "block_type": "paragraph",
                "text": "Pendingquartz policy wording.",
            }],
            "warnings": [],
        }
        knowledge_store.save_structured_document(
            structured, filename="pending.txt"
        )
        self.assertEqual(knowledge_store.search_chunks("Pendingquartz"), [])

        saved = knowledge_store.save_text(
            "Approved note",
            "Approvedcobalt knowledge wording.",
            source_type="Docs",
        )
        self.assertTrue(knowledge_store.search_chunks("Approvedcobalt"))
        self.assertEqual(saved["source_type"], "Docs")


if __name__ == "__main__":
    unittest.main()
