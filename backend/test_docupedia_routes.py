import asyncio
import copy
import os
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

from docx import Document
from fastapi import FastAPI, UploadFile
from fastapi.testclient import TestClient

import docupedia_routes
import document_extractor
import knowledge_routes
import knowledge_store
from test_document_upload_routes import GeneratedStream, _docx_bytes, _pdf_bytes


BASE = "/api/knowledge/docupedia"
RAW = "# Intake\nRawquartz source wording.\n\nKeepcobalt unchanged instruction."


class DocupediaRouteTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        database = Path(self.directory.name) / "knowledge.db"
        database_patch = patch.object(knowledge_store, "DATABASE_PATH", database)
        database_patch.start()
        self.addCleanup(database_patch.stop)
        secret_patch = patch.dict(os.environ, {
            "SESSION_SECRET": "docupedia-test-session-secret-with-at-least-32-characters",
        })
        secret_patch.start()
        self.addCleanup(secret_patch.stop)
        knowledge_store.initialize_database()
        self.app = FastAPI()
        self.app.include_router(knowledge_routes.router)
        self.app.include_router(docupedia_routes.router)
        self.client = TestClient(self.app)
        self.addCleanup(self.client.close)
        self.admin = {
            "Authorization": "Bearer " + knowledge_routes.create_access_token(
                "admin@example.test", "Admin"
            )
        }
        self.employee = {
            "Authorization": "Bearer " + knowledge_routes.create_access_token(
                "employee@example.test", "Employee"
            )
        }
        path = Path(self.directory.name) / "procedure.txt"
        path.write_bytes(RAW.encode("utf-8"))
        self.structured = document_extractor.extract_structured_document(path)

    def save_payload(self, **overrides):
        return {
            "filename": "procedure.txt",
            "structured_document": copy.deepcopy(self.structured),
            "idempotency_key": str(uuid4()),
            **overrides,
        }

    def save(self, payload=None):
        response = self.client.post(
            BASE + "/save", json=payload or self.save_payload(), headers=self.admin,
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["source"]

    def approve(self, source_id, text, key=None):
        return self.client.post(
            f"{BASE}/{source_id}/approve",
            json={"text": text, "idempotency_key": key or str(uuid4())},
            headers=self.admin,
        )

    def table_counts(self):
        with knowledge_store._connection() as connection:
            return {
                table: connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                for table in (
                    "knowledge_sources", "knowledge_chunks", "document_records",
                    "document_blocks", "document_block_cells", "document_sections",
                    "knowledge_chunk_blocks", "knowledge_chunk_sections",
                )
            }

    def test_all_endpoints_require_admin(self):
        source_id = self.save()["id"]
        for headers, status in (({}, 401), (self.employee, 403)):
            with self.subTest(status=status):
                requests = [
                    ("post", BASE + "/preview", {
                        "files": [("files", ("procedure.txt", RAW.encode(), "text/plain"))],
                    }),
                    ("post", BASE + "/save", {"json": self.save_payload()}),
                    ("get", f"{BASE}/{source_id}", {}),
                    ("post", f"{BASE}/{source_id}/approve", {
                        "json": {"text": "Approved wording.", "idempotency_key": str(uuid4())},
                    }),
                ]
                for method, url, kwargs in requests:
                    response = self.client.request(method, url, headers=headers, **kwargs)
                    self.assertEqual(response.status_code, status, response.text)

    def test_main_registers_all_docupedia_routes(self):
        import main

        schema = main.app.openapi()["paths"]
        for path, method in (
            (BASE + "/preview", "post"),
            (BASE + "/save", "post"),
            (BASE + "/{source_id}", "get"),
            (BASE + "/{source_id}/approve", "post"),
        ):
            self.assertIn(method, schema[path])

    def test_preview_supported_modern_formats_without_any_persistence(self):
        before = self.table_counts()
        with patch.object(
            knowledge_store, "initialize_database", side_effect=AssertionError("No persistence")
        ), patch.object(
            knowledge_store, "save_structured_document"
        ) as save, patch.object(knowledge_store, "save_text") as save_text:
            response = self.client.post(
                BASE + "/preview", headers=self.admin,
                files=[
                    ("files", ("procedure.pdf", _pdf_bytes(), "application/pdf")),
                    ("files", ("procedure.docx", _docx_bytes(), "application/octet-stream")),
                    ("files", ("procedure.txt", RAW.encode("utf-8"), "text/plain")),
                ],
            )
        self.assertEqual(response.status_code, 200, response.text)
        results = response.json()["results"]
        self.assertEqual([result["status"] for result in results], ["success"] * 3)
        self.assertEqual(
            [result["structured_document"]["format"] for result in results],
            ["pdf", "docx", "txt"],
        )
        self.assertEqual(results[2]["structured_document"], self.structured)
        self.assertEqual(results[0]["structured_document"]["blocks"][0]["page_start"], 1)
        self.assertTrue(any(
            block["block_type"] == "table_row"
            for block in results[1]["structured_document"]["blocks"]
        ))
        self.assertIn("warnings", results[1]["structured_document"])
        save.assert_not_called()
        save_text.assert_not_called()
        self.assertEqual(self.table_counts(), before)

    def test_doc_preview_reuses_conversion_and_reports_missing_libreoffice(self):
        def fake_conversion(command, **_kwargs):
            output = Path(command[command.index("--outdir") + 1]) / "document.docx"
            document = Document()
            document.add_heading("Legacy procedure", level=1)
            document.add_paragraph("Retain legacy business wording.")
            document.save(output)
            return type("Completed", (), {"returncode": 0, "stdout": "", "stderr": ""})()

        with patch.object(
            document_extractor, "_soffice_executable", return_value="soffice"
        ), patch.object(document_extractor.subprocess, "run", side_effect=fake_conversion):
            result = self.client.post(
                BASE + "/preview", headers=self.admin,
                files=[("files", ("procedure.doc", b"legacy input"))],
            ).json()["results"][0]
        self.assertEqual(result["status"], "success", result)
        self.assertEqual(result["structured_document"]["format"], "doc")
        self.assertIn("legacy_format_converted", {
            warning["code"] for warning in result["structured_document"]["warnings"]
        })
        with patch.object(document_extractor, "_soffice_executable", return_value=None):
            result = self.client.post(
                BASE + "/preview", headers=self.admin,
                files=[("files", ("procedure.doc", b"legacy input"))],
            ).json()["results"][0]
        self.assertEqual(result["status"], "error")
        self.assertIn("Legacy DOC files require LibreOffice", result["error"])

    def test_preview_validation_and_mixed_failures(self):
        response = self.client.post(
            BASE + "/preview", headers=self.admin,
            files=[
                ("files", ("good.txt", RAW.encode())),
                ("files", ("bad.xlsx", b"spreadsheet")),
                ("files", ("broken.pdf", b"not a pdf")),
                ("files", ("empty.txt", b"")),
                ("files", ("../bad.txt", b"bad name")),
                ("files", ("binary.txt", b"\x00\x01")),
            ],
        )
        self.assertEqual(response.status_code, 200)
        results = response.json()["results"]
        self.assertEqual([result["status"] for result in results], ["success"] + ["error"] * 5)
        self.assertIn("PDF, DOCX, DOC and TXT only", results[1]["error"])
        self.assertIn("PDF", results[2]["error"])
        self.assertIn("empty", results[3]["error"])
        self.assertIn("filename", results[4]["error"])
        self.assertIn("binary", results[5]["error"])
        self.assertEqual(
            self.client.post(BASE + "/preview", headers=self.admin).status_code, 422
        )

    def test_streaming_limit_and_temporary_cleanup(self):
        paths = []

        def extract(path):
            paths.append(path)
            self.assertTrue(path.is_file())
            return self.structured

        limit = knowledge_routes.MAX_DOCUMENT_FILE_BYTES
        uploads = [
            UploadFile(filename="boundary.txt", file=GeneratedStream(limit)),
            UploadFile(filename="oversize.txt", file=GeneratedStream(limit + 1)),
        ]
        with patch.object(docupedia_routes, "extract_structured_document", side_effect=extract):
            results = asyncio.run(docupedia_routes.preview_docupedia(
                files=uploads, _user={"role": "Admin"}
            ))["results"]
        self.assertEqual(results[0]["status"], "success")
        self.assertEqual(results[0]["metadata"]["size_bytes"], 50 * 1024 * 1024)
        self.assertEqual(results[1]["status"], "error")
        self.assertIn("50 MiB", results[1]["error"])
        self.assertTrue(all(upload.file.closed for upload in uploads))
        self.assertTrue(all(not path.parent.exists() for path in paths))

    def test_preview_cleanup_on_extractor_failure_and_response_size_policy(self):
        paths = []

        def fail(path):
            paths.append(path)
            raise document_extractor.DocumentExtractionError("Malformed content.")

        with patch.object(docupedia_routes, "extract_structured_document", side_effect=fail):
            response = self.client.post(
                BASE + "/preview", headers=self.admin,
                files=[("files", ("failure.txt", b"input"))],
            )
        self.assertEqual(response.json()["results"][0]["status"], "error")
        self.assertFalse(paths[0].parent.exists())
        with patch.object(docupedia_routes, "MAX_DOCUMENT_PREVIEW_RESPONSE_BYTES", 100):
            result = self.client.post(
                BASE + "/preview", headers=self.admin,
                files=[("files", ("large.txt", RAW.encode()))],
            ).json()["results"][0]
        self.assertEqual(result["status"], "error")
        self.assertIsNone(result["structured_document"])
        self.assertIn("16 MiB", result["error"])

    def test_pending_save_and_structured_get_preserve_raw_and_hide_employee_content(self):
        source = self.save()
        response = self.client.get(f"{BASE}/{source['id']}", headers=self.admin)
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["approval_status"], "pending")
        self.assertEqual(body["source"]["original_text"], "")
        self.assertEqual(body["source"]["filename"], "procedure.txt")
        self.assertEqual(body["structured_document"], self.structured)
        employee = self.client.post(
            "/api/knowledge/interview-context", json={"query": "Rawquartz"},
            headers=self.employee,
        )
        self.assertEqual(employee.status_code, 404)
        with knowledge_store._connection() as connection:
            self.assertGreater(connection.execute(
                "SELECT COUNT(*) FROM knowledge_chunks WHERE source_id = ?", (source["id"],)
            ).fetchone()[0], 0)

    def test_docx_preview_save_and_approval_preserve_cells_and_raw_provenance(self):
        preview = self.client.post(
            BASE + "/preview", headers=self.admin,
            files=[("files", ("procedure.docx", _docx_bytes()))],
        ).json()["results"][0]["structured_document"]
        source_id = self.save(self.save_payload(
            filename="procedure.docx", structured_document=preview,
        ))["id"]
        result = self.approve(source_id, preview["original_text"])
        self.assertEqual(result.status_code, 200, result.text)
        retrieved = self.client.get(f"{BASE}/{source_id}", headers=self.admin).json()
        self.assertEqual(retrieved["structured_document"], preview)
        self.assertEqual(retrieved["approval_status"], "approved")
        with knowledge_store._connection() as connection:
            row_links = connection.execute(
                """
                SELECT b.block_type FROM knowledge_chunk_blocks AS cb
                JOIN document_blocks AS b ON b.id = cb.block_id
                WHERE cb.source_id = ?
                """,
                (source_id,),
            ).fetchall()
        self.assertIn("table_row", [row["block_type"] for row in row_links])
        self.assertIn("table", [row["block_type"] for row in row_links])

    def test_explicit_pending_status_blocks_search_even_with_nonempty_canonical_text(self):
        source_id = self.save()["id"]
        with knowledge_store._connection() as connection:
            connection.execute(
                "UPDATE knowledge_sources SET original_text = ? WHERE id = ?",
                ("Unapproved staged text.", source_id),
            )
        self.assertEqual(knowledge_store.search_chunks("Rawquartz"), [])

    def test_save_duplicate_retries_and_changed_payload_conflict(self):
        payload = self.save_payload()
        first = self.save(payload)
        before = self.table_counts()
        second = self.save(payload)
        self.assertEqual(first["id"], second["id"])
        self.assertEqual(self.table_counts(), before)
        payload["structured_document"]["blocks"][1]["text"] = "Changed wording."
        response = self.client.post(BASE + "/save", json=payload, headers=self.admin)
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.table_counts(), before)

    def test_concurrent_save_retries_create_one_complete_source(self):
        payload = self.save_payload()
        barrier = threading.Barrier(2)

        def submit():
            with TestClient(self.app) as client:
                barrier.wait(timeout=10)
                return client.post(BASE + "/save", json=payload, headers=self.admin)

        with ThreadPoolExecutor(max_workers=2) as executor:
            responses = list(executor.map(lambda _index: submit(), range(2)))
        self.assertEqual([response.status_code for response in responses], [200, 200])
        ids = {response.json()["source"]["id"] for response in responses}
        self.assertEqual(len(ids), 1)
        self.assertEqual(self.table_counts()["knowledge_sources"], 1)
        self.assertEqual(
            knowledge_store.get_structured_document(ids.pop()), self.structured
        )

    def test_save_validates_filename_format_and_structured_relations(self):
        payloads = [
            self.save_payload(filename="../escape.txt"),
            self.save_payload(filename="wrong.pdf"),
            self.save_payload(idempotency_key="invalid"),
        ]
        missing_key = self.save_payload()
        del missing_key["idempotency_key"]
        payloads.append(missing_key)
        for edit in (
            lambda doc: doc.update(original_text=" "),
            lambda doc: doc.update(blocks=[]),
            lambda doc: doc["blocks"][1].update(ordinal=0),
            lambda doc: doc["blocks"][1].update(page_start=-1),
            lambda doc: doc["blocks"][1].update(parent_block_ordinal=99),
            lambda doc: doc["blocks"][2].update(parent_heading_block_ordinal=1),
            lambda doc: doc["blocks"][1].update(cells=[{"cell_index": 0, "text": None}]),
        ):
            payload = self.save_payload()
            edit(payload["structured_document"])
            payloads.append(payload)
        for payload in payloads:
            with self.subTest(payload=payload):
                response = self.client.post(BASE + "/save", json=payload, headers=self.admin)
                self.assertEqual(response.status_code, 422, response.text)
        self.assertEqual(self.table_counts()["knowledge_sources"], 0)

    def test_save_explicit_size_policy_rejects_without_truncating(self):
        with patch.object(docupedia_routes, "MAX_DOCUMENT_PREVIEW_RESPONSE_BYTES", 100):
            response = self.client.post(
                BASE + "/save", json=self.save_payload(), headers=self.admin
            )
        self.assertEqual(response.status_code, 413)
        self.assertEqual(self.table_counts()["document_records"], 0)

    def test_approval_publishes_only_edited_text_and_preserves_matching_provenance(self):
        source = self.save()
        source_id = source["id"]
        approved = "Approvedsapphire revised policy.\n\nKeepcobalt unchanged instruction."
        response = self.approve(source_id, approved)
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        self.assertEqual(result["approval_status"], "approved")
        self.assertEqual(result["source"]["original_text"], approved)
        fetched = self.client.get(f"{BASE}/{source_id}", headers=self.admin).json()
        self.assertEqual(fetched["structured_document"], self.structured)
        self.assertEqual(fetched["approval_status"], "approved")
        self.assertEqual(knowledge_store.search_chunks("Rawquartz"), [])
        employee = self.client.post(
            "/api/knowledge/interview-context", json={"query": "Approvedsapphire"},
            headers=self.employee,
        )
        self.assertEqual(employee.status_code, 200, employee.text)
        self.assertEqual(employee.json()["chunks"][0]["content"], approved)
        with knowledge_store._connection() as connection:
            links = connection.execute(
                """
                SELECT b.text FROM knowledge_chunk_blocks AS cb
                JOIN document_blocks AS b ON b.id = cb.block_id
                WHERE cb.source_id = ? AND cb.start_offset IS NOT NULL
                """,
                (source_id,),
            ).fetchall()
            self.assertEqual([row["text"] for row in links], ["Keepcobalt unchanged instruction."])
            self.assertGreater(connection.execute(
                "SELECT COUNT(*) FROM knowledge_chunk_sections WHERE source_id = ?", (source_id,)
            ).fetchone()[0], 0)
            self.assertEqual(connection.execute("PRAGMA foreign_key_check").fetchall(), [])

    def test_approval_retry_is_noop_and_conflicting_approvals_rejected(self):
        payload = self.save_payload()
        source_id = self.save(payload)["id"]
        key = str(uuid4())
        first = self.approve(source_id, "Approved wording.", key)
        self.assertEqual(first.status_code, 200)
        before = self.table_counts()
        retry = self.approve(source_id, "Approved wording.", key)
        self.assertEqual(retry.json(), first.json())
        self.assertEqual(self.table_counts(), before)
        for text, request_key in (
            ("Different wording.", key), ("Approved wording.", str(uuid4())),
        ):
            self.assertEqual(self.approve(source_id, text, request_key).status_code, 409)
        self.assertEqual(self.save(payload)["id"], source_id)
        self.assertEqual(
            knowledge_store.get_source(source_id)["original_text"], "Approved wording."
        )

    def test_concurrent_conflicting_approvals_have_one_winner(self):
        source_id = self.save()["id"]
        barrier = threading.Barrier(2)

        def approve(text):
            with TestClient(self.app) as client:
                barrier.wait(timeout=10)
                return client.post(
                    f"{BASE}/{source_id}/approve", headers=self.admin,
                    json={"text": text, "idempotency_key": str(uuid4())},
                )

        with ThreadPoolExecutor(max_workers=2) as executor:
            responses = list(executor.map(approve, ["Firstapproved wording.", "Secondapproved wording."]))
        self.assertEqual(sorted(response.status_code for response in responses), [200, 409])
        winner = next(response for response in responses if response.status_code == 200)
        stored = knowledge_store.get_source(source_id)
        self.assertEqual(stored["original_text"], winner.json()["source"]["original_text"])
        with knowledge_store._connection() as connection:
            chunks = connection.execute(
                "SELECT content FROM knowledge_chunks WHERE source_id = ?", (source_id,)
            ).fetchall()
        self.assertEqual([row["content"] for row in chunks], [stored["original_text"]])

    def test_approval_rejects_empty_text_and_unknown_or_non_docupedia_source(self):
        source_id = self.save()["id"]
        self.assertEqual(self.approve(source_id, "  ").status_code, 422)
        self.assertEqual(self.approve(source_id, "").status_code, 422)
        self.assertEqual(self.approve(source_id, "x" * 200_001).status_code, 422)
        normal = knowledge_store.save_text("Normal", "Ordinary approved source.")
        for unknown in (str(uuid4()), normal["id"]):
            self.assertEqual(self.approve(unknown, "Reviewed wording.").status_code, 404)
            self.assertEqual(
                self.client.get(f"{BASE}/{unknown}", headers=self.admin).status_code, 404
            )

    def test_pending_source_cannot_bypass_approval_via_generic_update(self):
        source_id = self.save()["id"]
        response = self.client.put(
            f"/api/knowledge/sources/{source_id}", headers=self.admin,
            json={"title": "Bypass", "text": "Bypassquartz wording.", "notes": ""},
        )
        self.assertEqual(response.status_code, 409)
        self.assertEqual(knowledge_store.search_chunks("Bypassquartz"), [])
        self.assertEqual(knowledge_store.get_source(source_id)["original_text"], "")

    def test_save_transaction_rolls_back_after_chunk_link_failure(self):
        with knowledge_store._connection() as connection:
            connection.execute("""
                CREATE TRIGGER reject_document_link BEFORE INSERT ON knowledge_chunk_blocks
                BEGIN SELECT RAISE(ABORT, 'forced link failure'); END
            """)
        before = self.table_counts()
        with self.assertLogs("docupedia_routes", level="ERROR"):
            response = self.client.post(
                BASE + "/save", json=self.save_payload(), headers=self.admin
            )
        self.assertEqual(response.status_code, 500)
        self.assertEqual(self.table_counts(), before)

    def test_approval_transaction_rolls_back_content_status_chunks_and_links(self):
        source_id = self.save()["id"]
        before = self.table_counts()
        with knowledge_store._connection() as connection:
            original_chunks = [tuple(row) for row in connection.execute(
                "SELECT * FROM knowledge_chunks WHERE source_id = ?", (source_id,)
            )]
            connection.execute("""
                CREATE TRIGGER reject_approved_link BEFORE INSERT ON knowledge_chunk_blocks
                BEGIN SELECT RAISE(ABORT, 'forced approval failure'); END
            """)
        with self.assertLogs("docupedia_routes", level="ERROR"):
            response = self.approve(
                source_id, "Approvedsapphire wording.\n\nKeepcobalt unchanged instruction."
            )
        self.assertEqual(response.status_code, 500)
        self.assertEqual(self.table_counts(), before)
        stored = knowledge_store.get_source(source_id)
        self.assertEqual(stored["metadata"]["docupedia"]["approval_status"], "pending")
        self.assertEqual(stored["original_text"], "")
        self.assertEqual(knowledge_store.get_structured_document(source_id), self.structured)
        self.assertEqual(knowledge_store.search_chunks("Approvedsapphire"), [])
        with knowledge_store._connection() as connection:
            self.assertEqual(
                [tuple(row) for row in connection.execute(
                    "SELECT * FROM knowledge_chunks WHERE source_id = ?", (source_id,)
                )],
                original_chunks,
            )
            connection.execute("DROP TRIGGER reject_approved_link")
        self.assertEqual(self.approve(source_id, "Approvedsapphire wording.").status_code, 200)


if __name__ == "__main__":
    unittest.main()
