import asyncio
import io
import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi import FastAPI, UploadFile
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from pptx import Presentation

import knowledge_routes
from knowledge_routes import create_access_token, router


def _pdf_bytes():
    writer = PdfWriter()
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    font_ref = writer._add_object(font)
    for content in ("PDF page one: Operations", "PDF page two: pressure 2.5 bar"):
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})
        })
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 14 Tf 50 740 Td ({content}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _docx_bytes():
    document = Document()
    document.add_heading("Process: café start-up", level=1)
    document.add_paragraph("Verify temperature is 2.5 °C.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Step"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Inspect valve"
    table.cell(1, 1).text = "Zoë"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_bytes():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Café checks"
    worksheet.append(["Step", "Value"])
    worksheet.append(["Pressure", 2.5])
    workbook.create_sheet("Owners").append(["Zoë", "Ready ✓"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pptx_bytes():
    presentation = Presentation()
    for title, body in (
        ("Start-up", "Check pressure: 2.5 bar."),
        ("Café shutdown", "Owner: Zoë."),
    ):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


class GeneratedStream:
    def __init__(self, length):
        self.remaining = length
        self.closed = False

    def read(self, size=-1):
        if self.remaining <= 0:
            return b""
        count = self.remaining if size < 0 else min(size, self.remaining)
        self.remaining -= count
        return b"x" * count

    def close(self):
        self.closed = True


class DocumentUploadRouteTests(unittest.TestCase):
    def setUp(self):
        self.secret = patch.dict(
            os.environ,
            {"SESSION_SECRET": "document-route-test-secret-at-least-32-chars"},
        )
        self.secret.start()
        self.addCleanup(self.secret.stop)
        app = FastAPI()
        app.include_router(router)
        self.client = TestClient(app)
        self.admin_token = create_access_token("admin@example.test", "Admin")
        self.employee_token = create_access_token("employee@example.test", "Employee")

    def upload(self, filename, content, token=None):
        headers = {}
        if token:
            headers["Authorization"] = "Bearer " + token
        return self.client.post(
            "/api/knowledge/documents/preview",
            files=[("files", (filename, content, "application/octet-stream"))],
            headers=headers,
        )

    def test_requires_admin_authorization(self):
        content = b"Preview text"
        self.assertEqual(self.upload("note.txt", content).status_code, 401)
        self.assertEqual(
            self.upload("note.txt", content, self.employee_token).status_code, 403
        )
        response = self.upload("note.txt", content, self.admin_token)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["results"][0]["status"], "success")

    def test_extracts_pdf_docx_xlsx_pptx_and_txt(self):
        samples = [
            ("operations.pdf", _pdf_bytes(), ["--- Page 1 ---", "--- Page 2 ---", "pressure 2.5 bar"]),
            ("operations.docx", _docx_bytes(), ["Process: café start-up", "Step | Owner", "Inspect valve | Zoë"]),
            ("operations.xlsx", _xlsx_bytes(), ["--- Sheet: Café checks ---", "--- Sheet: Owners ---", "Pressure\t2.5", "Zoë\tReady ✓"]),
            ("operations.pptx", _pptx_bytes(), ["--- Slide 1 ---", "--- Slide 2 ---", "Title: Start-up", "Owner: Zoë"]),
            ("operations.txt", "Helix café – Zoë ✓".encode("utf-8"), ["Helix café – Zoë ✓"]),
        ]
        with patch.object(knowledge_routes.knowledge_store, "save_text") as save_text:
            response = self.client.post(
                "/api/knowledge/documents/preview",
                files=[
                    ("files", (filename, content, "application/octet-stream"))
                    for filename, content, _expected in samples
                ],
                headers={"Authorization": f"Bearer {self.admin_token}"},
            )
        save_text.assert_not_called()
        self.assertEqual(response.status_code, 200)
        results = response.json()["results"]
        self.assertEqual(len(results), len(samples))
        for result, (filename, _content, expected) in zip(results, samples):
            with self.subTest(filename=filename):
                self.assertEqual(result["status"], "success")
                self.assertEqual(result["filename"], filename)
                for fragment in expected:
                    self.assertIn(fragment, result["text"])
                self.assertGreater(result["metadata"]["size_bytes"], 0)

    def test_reports_unsupported_malformed_empty_and_mixed_uploads(self):
        mixed = self.client.post(
            "/api/knowledge/documents/preview",
            files=[
                ("files", ("good.txt", b"Known process text", "text/plain")),
                ("files", ("unsupported.odt", b"data", "application/octet-stream")),
                ("files", ("broken.xlsx", b"not an xlsx archive", "application/octet-stream")),
                ("files", ("empty.txt", b"", "text/plain")),
            ],
            headers={"Authorization": f"Bearer {self.admin_token}"},
        )
        self.assertEqual(mixed.status_code, 200)
        results = mixed.json()["results"]
        self.assertEqual([result["status"] for result in results], [
            "success", "error", "error", "error",
        ])
        self.assertIn("Known process text", results[0]["text"])
        self.assertIn("Unsupported document extension", results[1]["error"])
        self.assertIn("Unable to read the XLSX", results[2]["error"])
        self.assertIn("empty", results[3]["error"])

    def test_rejects_malformed_filename(self):
        result = self.upload("../escape.txt", b"content", self.admin_token).json()["results"][0]
        self.assertEqual(result["status"], "error")
        self.assertIn("Invalid document filename", result["error"])

    def test_exact_limit_is_accepted_and_one_byte_over_is_rejected(self):
        exact_stream = GeneratedStream(knowledge_routes.MAX_DOCUMENT_FILE_BYTES)
        exact_upload = UploadFile(
            filename="exact.txt", file=exact_stream, size=None
        )
        with patch.object(knowledge_routes, "extract_document", return_value="Boundary content"):
            exact_result = asyncio.run(
                knowledge_routes.preview_documents([exact_upload], {"role": "Admin"})
            )["results"][0]
        self.assertEqual(exact_result["status"], "success")
        self.assertEqual(
            exact_result["metadata"]["size_bytes"],
            knowledge_routes.MAX_DOCUMENT_FILE_BYTES,
        )
        self.assertTrue(exact_stream.closed)

        oversized_stream = GeneratedStream(knowledge_routes.MAX_DOCUMENT_FILE_BYTES + 1)
        oversized_upload = UploadFile(
            filename="oversized.txt", file=oversized_stream, size=None
        )
        with patch.object(knowledge_routes, "extract_document") as extractor:
            oversized_result = asyncio.run(
                knowledge_routes.preview_documents([oversized_upload], {"role": "Admin"})
            )["results"][0]
        self.assertEqual(oversized_result["status"], "error")
        self.assertIn("50 MiB", oversized_result["error"])
        self.assertTrue(oversized_stream.closed)
        extractor.assert_not_called()

    def test_temporary_upload_directories_are_cleaned_up(self):
        original_factory = tempfile.TemporaryDirectory
        created_directories = []

        def capture_directory(*args, **kwargs):
            temporary_directory = original_factory(*args, **kwargs)
            created_directories.append(Path(temporary_directory.name))
            return temporary_directory

        with patch.object(
            knowledge_routes,
            "tempfile",
            wraps=knowledge_routes.tempfile,
        ) as tempfile_module:
            tempfile_module.TemporaryDirectory.side_effect = capture_directory
            response = self.upload("cleanup.txt", b"Temporary content", self.admin_token)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(created_directories), 1)
        self.assertFalse(created_directories[0].exists())

    def test_preview_text_exceeding_response_budget_is_reported_not_truncated(self):
        oversized_preview = "x" * (
            knowledge_routes.MAX_DOCUMENT_PREVIEW_RESPONSE_BYTES + 1
        )
        with patch.object(
            knowledge_routes, "extract_document", return_value=oversized_preview
        ):
            result = self.upload(
                "large-preview.txt", b"source", self.admin_token
            ).json()["results"][0]
        self.assertEqual(result["status"], "error")
        self.assertIsNone(result["text"])
        self.assertIn("16 MiB limit", result["error"])


class DocumentSaveRouteTests(unittest.TestCase):
    def setUp(self):
        self.secret = patch.dict(
            os.environ,
            {"SESSION_SECRET": "document-save-test-secret-at-least-32-chars"},
        )
        self.secret.start()
        self.addCleanup(self.secret.stop)
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.original_database = knowledge_routes.knowledge_store.DATABASE_PATH
        knowledge_routes.knowledge_store.DATABASE_PATH = (
            Path(self.directory.name) / "knowledge.db"
        )
        self.addCleanup(
            setattr,
            knowledge_routes.knowledge_store,
            "DATABASE_PATH",
            self.original_database,
        )
        app = FastAPI()
        app.include_router(router)
        self.client = TestClient(app)
        self.admin_token = create_access_token("admin@example.test", "Admin")
        self.employee_token = create_access_token("employee@example.test", "Employee")

    def save(self, documents, token=None):
        headers = {}
        if token:
            headers["Authorization"] = "Bearer " + token
        return self.client.post(
            "/api/knowledge/documents/save",
            json={"documents": documents},
            headers=headers,
        )

    @staticmethod
    def document(filename, text, key=None):
        from uuid import uuid4

        return {
            "filename": filename,
            "text": text,
            "metadata": {"format": filename.rsplit(".", 1)[-1], "size_bytes": 128},
            "idempotency_key": str(key or uuid4()),
        }

    def test_requires_admin_authorization(self):
        documents = [self.document("manual.txt", "Admin-only content.")]
        self.assertEqual(self.save(documents).status_code, 401)
        self.assertEqual(
            self.save(documents, self.employee_token).status_code, 403
        )
        self.assertEqual(
            self.save(documents, self.admin_token).status_code, 200
        )

    def test_persists_document_metadata_chunks_and_search_retrieval(self):
        text = "DocSearchToken describes the reviewed valve inspection procedure."
        response = self.save(
            [self.document("valve-guide.pdf", text)],
            self.admin_token,
        )
        self.assertEqual(response.status_code, 200)
        result = response.json()["results"][0]
        self.assertEqual(result["status"], "success")
        source = knowledge_routes.knowledge_store.get_source(result["source"]["id"])
        self.assertEqual(source["source_type"], "Document")
        self.assertEqual(source["filename"], "valve-guide.pdf")
        self.assertEqual(source["title"], "valve-guide.pdf")
        self.assertEqual(source["original_text"], text)
        self.assertEqual(source["metadata"]["original_filename"], "valve-guide.pdf")
        self.assertEqual(source["metadata"]["format"], "pdf")
        chunks = knowledge_routes.knowledge_store.search_chunks("DocSearchToken")
        self.assertTrue(any(chunk["source_id"] == source["id"] for chunk in chunks))

    def test_saves_edited_preview_text(self):
        document = self.document(
            "edited.docx",
            "AdminEditedToken: use the updated pressure threshold.",
        )
        result = self.save([document], self.admin_token).json()["results"][0]
        source = knowledge_routes.knowledge_store.get_source(result["source"]["id"])
        self.assertEqual(
            source["original_text"],
            "AdminEditedToken: use the updated pressure threshold.",
        )
        self.assertTrue(
            knowledge_routes.knowledge_store.search_chunks("AdminEditedToken")
        )

    def test_saves_multiple_documents_independently_and_reports_invalid_content(self):
        response = self.save(
            [
                self.document("first.txt", "FirstDocToken is saved."),
                self.document("blank.txt", "   "),
                self.document("second.xlsx", "SecondDocToken is saved."),
                self.document("unsupported.odt", "UnsupportedDocToken."),
            ],
            self.admin_token,
        )
        self.assertEqual(response.status_code, 200)
        results = response.json()["results"]
        self.assertEqual(
            [result["status"] for result in results],
            ["success", "error", "success", "error"],
        )
        self.assertIn("must not be blank", results[1]["error"])
        self.assertIn("Unsupported document extension", results[3]["error"])
        self.assertEqual(len(knowledge_routes.knowledge_store.list_sources()), 2)

    def test_idempotent_retry_does_not_create_duplicate_source(self):
        document = self.document("retry.txt", "RetryDocToken is stored once.")
        first = self.save([document], self.admin_token).json()["results"][0]
        retry = self.save([document], self.admin_token).json()["results"][0]
        self.assertEqual(first["status"], "success")
        self.assertEqual(retry["status"], "success")
        self.assertEqual(first["source"]["id"], retry["source"]["id"])
        self.assertEqual(len(knowledge_routes.knowledge_store.list_sources()), 1)
        changed_retry = {**document, "text": "ChangedRetryToken would duplicate content."}
        changed_result = self.save(
            [changed_retry], self.admin_token
        ).json()["results"][0]
        self.assertEqual(changed_result["status"], "error")
        self.assertIn("already used with different content", changed_result["error"])
        self.assertEqual(len(knowledge_routes.knowledge_store.list_sources()), 1)

    def test_integrity_error_retry_returns_existing_source_with_metadata(self):
        store = knowledge_routes.knowledge_store
        key = "integrity-fallback-document-key"
        metadata = {"format": "txt", "original_filename": "retry.txt"}
        text = "IntegrityFallbackToken stored exactly once."
        original_connection = store._connection

        original_source = store.save_text(
            "retry.txt",
            text,
            source_type="Document",
            filename="retry.txt",
            metadata=metadata,
            idempotency_key=key,
        )

        class ConcurrentInsertConflict:
            def __enter__(self):
                raise sqlite3.IntegrityError("unique idempotency key conflict")

            def __exit__(self, _exception_type, _exception, _traceback):
                return False

        def conflict_on_insert():
            nonlocal connection_calls
            connection_calls += 1
            if connection_calls == 2:
                return ConcurrentInsertConflict()
            return original_connection()

        connection_calls = 0
        with patch.object(store, "_connection", side_effect=conflict_on_insert):
            retried_source = store.save_text(
                "retry.txt",
                text,
                source_type="Document",
                filename="retry.txt",
                metadata=metadata,
                idempotency_key=key,
            )

        self.assertEqual(retried_source["id"], original_source["id"])
        self.assertEqual(retried_source["metadata"], metadata)
        self.assertEqual(len(store.list_sources()), 1)

        connection_calls = 0
        with patch.object(store, "_connection", side_effect=conflict_on_insert):
            with self.assertRaisesRegex(ValueError, "already used with different content"):
                store.save_text(
                    "retry.txt",
                    "ChangedIntegrityFallbackToken content.",
                    source_type="Document",
                    filename="retry.txt",
                    metadata=metadata,
                    idempotency_key=key,
                )
        self.assertEqual(len(store.list_sources()), 1)


if __name__ == "__main__":
    unittest.main()
