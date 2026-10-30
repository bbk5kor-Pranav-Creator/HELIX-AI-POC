"""Extract readable text from common office document formats."""

import importlib
import os
import re
import shutil
import subprocess
import tempfile
from datetime import date, datetime, time
from pathlib import Path


SUPPORTED_FORMATS = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".pptx", ".ppt", ".txt"}
LEGACY_CONVERSIONS = {
    ".doc": ".docx",
    ".xls": ".xlsx",
    ".ppt": ".pptx",
}
CONVERSION_TIMEOUT_SECONDS = 120
DOCPEDIA_FORMATS = {".pdf", ".docx", ".doc", ".txt"}


class DocumentExtractionError(ValueError):
    """A document is unsupported, invalid, or cannot provide readable text."""


def _require_package(module_name, package_name, file_format):
    try:
        return importlib.import_module(module_name)
    except ImportError as error:
        raise DocumentExtractionError(
            f"Extracting {file_format.upper()} files requires the "
            f"{package_name} package."
        ) from error


def _ensure_content(text, file_format):
    if not text.strip():
        raise DocumentExtractionError(
            f"No readable text or supported content was found in the "
            f"{file_format.upper()} file."
        )
    return text


def _extract_pdf(path):
    pdf_module = _require_package("pypdf", "pypdf", "pdf")
    try:
        reader = pdf_module.PdfReader(str(path), strict=True)
        if reader.is_encrypted:
            try:
                decrypted = reader.decrypt("")
            except Exception as error:
                raise DocumentExtractionError(
                    "The PDF is password-protected and cannot be opened."
                ) from error
            if not decrypted:
                raise DocumentExtractionError(
                    "The PDF is password-protected and cannot be opened."
                )

        pages = []
        has_content = False
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            has_content = has_content or bool(page_text.strip())
            pages.append(f"--- Page {page_number} ---\n{page_text.strip()}")
        if not has_content:
            raise DocumentExtractionError(
                "No readable text or supported content was found in the PDF file."
            )
        return _ensure_content("\n\n".join(pages), "pdf")
    except DocumentExtractionError:
        raise
    except Exception as error:
        raise DocumentExtractionError(
            "Unable to read the PDF; it may be corrupted, malformed, or "
            "password-protected."
        ) from error


def _word_body_content(document, paragraph_type, table_type, qn):
    content = []
    body = document.element.body
    for element in body.iterchildren():
        if element.tag == qn("w:p"):
            paragraph = paragraph_type(element, document)
            text = paragraph.text.strip()
            if text:
                content.append(text)
        elif element.tag == qn("w:tbl"):
            table = table_type(element, document)
            for row in table.rows:
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    content.append(" | ".join(cells))
    return content


def _extract_docx(path, source_format="docx"):
    docx_module = _require_package("docx", "python-docx", source_format)
    try:
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = docx_module.Document(str(path))
        content = _word_body_content(document, Paragraph, Table, qn)
        return _ensure_content("\n".join(content), source_format)
    except DocumentExtractionError:
        raise
    except Exception as error:
        raise DocumentExtractionError(
            f"Unable to read the {source_format.upper()}; it may be corrupted, malformed, or "
            "password-protected."
        ) from error


def _docx_heading_level(paragraph):
    style_name = getattr(paragraph.style, "name", "") or ""
    match = re.fullmatch(r"Heading\s+([1-9])", style_name, flags=re.IGNORECASE)
    if match:
        return int(match.group(1))

    properties = paragraph._p.pPr
    outline_level = properties.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}outlineLvl"
    ) if properties is not None else None
    if outline_level is not None:
        try:
            value = outline_level.get(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"
            )
            return int(value) + 1
        except (TypeError, ValueError):
            return None
    if style_name.lower() == "title":
        return 1
    return None


def _structured_block(ordinal, block_type, text="", **provenance):
    return {
        "ordinal": ordinal,
        "block_type": block_type,
        "text": text,
        **provenance,
    }


def _extract_docx_structured(path, source_format="docx"):
    docx_module = _require_package("docx", "python-docx", source_format)
    try:
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = docx_module.Document(str(path))
        blocks = []
        warnings = []
        paragraph_index = 0
        table_index = 0
        for element in document.element.body.iterchildren():
            if element.tag == qn("w:p"):
                paragraph = Paragraph(element, document)
                text = paragraph.text.strip()
                if text:
                    heading_level = _docx_heading_level(paragraph)
                    block_type = "heading" if heading_level is not None else "paragraph"
                    blocks.append(_structured_block(
                        len(blocks),
                        block_type,
                        text,
                        heading_level=heading_level,
                        paragraph_index=paragraph_index,
                        page_start=None,
                        page_end=None,
                        locator={"style": getattr(paragraph.style, "name", None)},
                    ))
                if element.xpath(".//w:drawing | .//w:pict | .//w:object"):
                    warnings.append({
                        "code": "embedded_content_not_extracted",
                        "message": (
                            "This paragraph contains an image or embedded object; "
                            "its visual or embedded content was not extracted."
                        ),
                        "paragraph_index": paragraph_index,
                    })
                paragraph_index += 1
            elif element.tag == qn("w:tbl"):
                table = Table(element, document)
                table_block = _structured_block(
                    len(blocks),
                    "table",
                    "",
                    table_index=table_index,
                    page_start=None,
                    page_end=None,
                    locator={},
                )
                blocks.append(table_block)
                for row_index, row in enumerate(table.rows):
                    cells = [
                        {
                            "cell_index": cell_index,
                            "column_index": cell_index,
                            "text": " ".join(cell.text.split()),
                        }
                        for cell_index, cell in enumerate(row.cells)
                    ]
                    row_text = " | ".join(cell["text"] for cell in cells)
                    blocks.append(_structured_block(
                        len(blocks),
                        "table_row",
                        row_text,
                        parent_block_ordinal=table_block["ordinal"],
                        table_index=table_index,
                        row_index=row_index,
                        page_start=None,
                        page_end=None,
                        cells=cells,
                        locator={},
                    ))
                if element.xpath(".//w:drawing | .//w:pict | .//w:object"):
                    warnings.append({
                        "code": "embedded_content_not_extracted",
                        "message": (
                            "This table contains an image or embedded object; "
                            "its visual or embedded content was not extracted."
                        ),
                        "table_index": table_index,
                    })
                table_index += 1

        raw_content = _word_body_content(
            document, Paragraph, Table, qn
        )
        original_text = "\n".join(raw_content)
        if not original_text.strip():
            raise DocumentExtractionError(
                f"No readable text or supported content was found in the "
                f"{source_format.upper()} file."
            )

        for section in document.sections:
            for part_name, part in (("header", section.header), ("footer", section.footer)):
                if any(paragraph.text.strip() for paragraph in part.paragraphs) or any(
                    any(cell.text.strip() for row in table.rows for cell in row.cells)
                    for table in part.tables
                ):
                    warnings.append({
                        "code": "document_part_not_extracted",
                        "message": f"Text in a DOCX {part_name} was not included.",
                    })

        return {
            "format": source_format.lstrip(".").lower(),
            "original_text": original_text,
            "blocks": blocks,
            "warnings": warnings,
        }
    except DocumentExtractionError:
        raise
    except Exception as error:
        raise DocumentExtractionError(
            f"Unable to read the {source_format.upper()}; it may be corrupted, malformed, or "
            "password-protected."
        ) from error


def _text_heading_level(text):
    candidate = text.strip()
    markdown_heading = re.match(r"^(#{1,6})\s+\S", candidate)
    if markdown_heading:
        return len(markdown_heading.group(1))

    numbered_heading = re.match(r"^(\d+(?:\.\d+)*)(?:[.)])?\s+\S", candidate)
    if numbered_heading and len(candidate) <= 120 and not candidate.endswith((".", "!", "?")):
        return numbered_heading.group(1).count(".") + 1

    if (
        len(candidate) <= 80
        and any(character.isalpha() for character in candidate)
        and candidate == candidate.upper()
        and not candidate.endswith((".", "!", "?"))
    ):
        return 1
    return None


def _page_contains_image(page):
    try:
        resources = page.get("/Resources")
        if resources is None:
            return False
        pending = [resources]
        visited = set()
        while pending:
            current = pending.pop()
            if hasattr(current, "get_object"):
                current = current.get_object()
            identity = id(current)
            if identity in visited:
                continue
            visited.add(identity)
            if hasattr(current, "get"):
                if current.get("/Subtype") == "/Image":
                    return True
                pending.extend(current.values())
            elif isinstance(current, (list, tuple)):
                pending.extend(current)
    except Exception:
        return False
    return False


def _extract_pdf_structured(path):
    pdf_module = _require_package("pypdf", "pypdf", "pdf")
    try:
        reader = pdf_module.PdfReader(str(path), strict=True)
        if reader.is_encrypted:
            try:
                decrypted = reader.decrypt("")
            except Exception as error:
                raise DocumentExtractionError(
                    "The PDF is password-protected and cannot be opened."
                ) from error
            if not decrypted:
                raise DocumentExtractionError(
                    "The PDF is password-protected and cannot be opened."
                )

        pages = []
        blocks = []
        warnings = []
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            pages.append(f"--- Page {page_number} ---\n{page_text.strip()}")
            has_images = _page_contains_image(page)
            meaningful_lines = [
                (line_index, line.strip())
                for line_index, line in enumerate(page_text.splitlines())
                if line.strip()
            ]
            if not meaningful_lines:
                warnings.append({
                    "code": "page_no_extractable_text",
                    "message": f"Page {page_number} contains no meaningful extractable text.",
                    "page_number": page_number,
                })
                if has_images:
                    warnings.append({
                        "code": "suspected_scanned_page",
                        "message": (
                            f"Page {page_number} contains images but no extractable text; "
                            "it may be scanned. OCR has not been run."
                        ),
                        "page_number": page_number,
                    })
                continue

            if has_images:
                warnings.append({
                    "code": "page_image_content_not_extracted",
                    "message": (
                        f"Page {page_number} contains image content that was not "
                        "processed with OCR."
                    ),
                    "page_number": page_number,
                })

            for line_index, line in meaningful_lines:
                heading_level = _text_heading_level(line)
                blocks.append(_structured_block(
                    len(blocks),
                    "heading" if heading_level is not None else "paragraph",
                    line,
                    heading_level=heading_level,
                    page_start=page_number,
                    page_end=page_number,
                    paragraph_index=None,
                    locator={"line_index": line_index},
                ))

        return {
            "format": "pdf",
            "original_text": "\n\n".join(pages),
            "blocks": blocks,
            "warnings": warnings,
        }
    except DocumentExtractionError:
        raise
    except Exception as error:
        raise DocumentExtractionError(
            "Unable to read the PDF; it may be corrupted, malformed, or "
            "password-protected."
        ) from error


def _soffice_executable():
    configured = os.getenv("SOFFICE_PATH")
    if configured:
        executable = shutil.which(configured)
        if executable:
            return executable
        if Path(configured).is_file():
            return configured
        return None
    return shutil.which("soffice") or shutil.which("libreoffice")


def _macro_disabled_profile(profile_path):
    user_profile = profile_path / "user"
    user_profile.mkdir()
    registry = user_profile / "registrymodifications.xcu"
    registry.write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
        'xmlns:xs="http://www.w3.org/2001/XMLSchema">\n'
        '  <item oor:path="/org.openoffice.Office.Common/Security/Scripting">\n'
        '    <prop oor:name="MacroSecurityLevel" oor:op="fuse">'
        '<value>3</value></prop>\n'
        '  </item>\n'
        '  <item oor:path="/org.openoffice.Office.Common/Load/Document">\n'
        '    <prop oor:name="UpdateDocMode" oor:op="fuse"><value>2</value></prop>\n'
        '  </item>\n'
        '</oor:items>\n',
        encoding="utf-8",
    )


def _convert_legacy(path, file_format):
    return _convert_legacy_impl(path, file_format, structured=False)


def _convert_legacy_structured(path, file_format):
    return _convert_legacy_impl(path, file_format, structured=True)


def _convert_legacy_impl(path, file_format, *, structured):
    executable = _soffice_executable()
    if not executable:
        raise DocumentExtractionError(
            f"Legacy {file_format.upper().lstrip('.')} files require LibreOffice "
            "(soffice) for conversion. Install LibreOffice or convert the "
            f"file to its modern format."
        )

    target_extension = LEGACY_CONVERSIONS[file_format]
    with tempfile.TemporaryDirectory(prefix="helix_document_conversion_") as temp:
        temp_path = Path(temp)
        output_path = temp_path / "converted"
        profile_path = temp_path / "profile"
        output_path.mkdir()
        profile_path.mkdir()
        _macro_disabled_profile(profile_path)
        profile_uri = profile_path.resolve().as_uri()
        command = [
            executable,
            f"-env:UserInstallation={profile_uri}",
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--norestore",
            "--convert-to",
            target_extension.lstrip("."),
            "--outdir",
            str(output_path),
            str(path),
        ]
        try:
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=CONVERSION_TIMEOUT_SECONDS,
                check=False,
            )
        except subprocess.TimeoutExpired as error:
            raise DocumentExtractionError(
                f"LibreOffice conversion of the {file_format.upper()} file "
                "timed out."
            ) from error
        except OSError as error:
            raise DocumentExtractionError(
                f"LibreOffice could not start to convert the "
                f"{file_format.upper()} file."
            ) from error

        converted_path = output_path / f"{path.stem}{target_extension}"
        if result.returncode != 0 or not converted_path.is_file():
            detail = (result.stderr or result.stdout).strip()
            suffix = f" Details: {detail}" if detail else ""
            raise DocumentExtractionError(
                f"LibreOffice could not convert the {file_format.upper()} "
                f"file. It may be corrupted, malformed, or password-protected."
                f"{suffix}"
            )

        if structured:
            if file_format != ".doc":
                raise DocumentExtractionError(
                    f"Structured extraction is not supported for {file_format.upper()}."
                )
            result = _extract_docx_structured(
                converted_path, source_format=file_format
            )
            result["warnings"].append({
                "code": "legacy_format_converted",
                "message": (
                    "DOC structure was read from a LibreOffice-converted DOCX; "
                    "conversion may alter formatting and exact source layout."
                ),
            })
            return result

        extractors = {
            ".doc": _extract_docx,
            ".xls": _extract_xlsx,
            ".ppt": _extract_pptx,
        }
        extractor = extractors[file_format]
        return extractor(converted_path, source_format=file_format)


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _extract_xlsx(path, source_format="xlsx"):
    openpyxl = _require_package("openpyxl", "openpyxl", source_format)
    workbook = None
    try:
        workbook = openpyxl.load_workbook(
            filename=path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        sheets = []
        has_content = False
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                cells = [_cell_text(value) for value in row]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    rows.append("\t".join(cells))
            has_content = has_content or bool(rows)
            sheet_content = f"--- Sheet: {worksheet.title} ---"
            if rows:
                sheet_content += "\n" + "\n".join(rows)
            sheets.append(sheet_content)
        if not has_content:
            raise DocumentExtractionError(
                "No readable text or supported content was found in the "
                f"{source_format.upper()} file."
            )
        return _ensure_content("\n\n".join(sheets), source_format)
    except DocumentExtractionError:
        raise
    except Exception as error:
        raise DocumentExtractionError(
            f"Unable to read the {source_format.upper()}; it may be corrupted, malformed, or "
            "password-protected."
        ) from error
    finally:
        if workbook is not None:
            workbook.close()


def _shape_text(shape, group_type):
    content = []
    if shape.shape_type == group_type:
        for child in shape.shapes:
            content.extend(_shape_text(child, group_type))
        return content
    if getattr(shape, "has_text_frame", False):
        text = shape.text.strip()
        if text:
            content.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                content.append(" | ".join(cells))
    return content


def _extract_pptx(path, source_format="pptx"):
    pptx_module = _require_package("pptx", "python-pptx", source_format)
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        presentation = pptx_module.Presentation(str(path))
        slides = []
        has_content = False
        for slide_number, slide in enumerate(presentation.slides, start=1):
            title_shape = slide.shapes.title
            title = title_shape.text.strip() if title_shape is not None else ""
            title_shape_id = title_shape.shape_id if title_shape is not None else None
            heading = f"--- Slide {slide_number} ---"
            if title:
                heading += f"\nTitle: {title}"
            body = []
            for shape in slide.shapes:
                if shape.shape_id == title_shape_id:
                    continue
                body.extend(_shape_text(shape, MSO_SHAPE_TYPE.GROUP))
            has_content = has_content or bool(title or body)
            slide_content = heading
            if body:
                slide_content += "\n" + "\n".join(body)
            slides.append(slide_content)
        if not has_content:
            raise DocumentExtractionError(
                "No readable text or supported content was found in the "
                f"{source_format.upper()} file."
            )
        return _ensure_content("\n\n".join(slides), source_format)
    except DocumentExtractionError:
        raise
    except Exception as error:
        raise DocumentExtractionError(
            f"Unable to read the {source_format.upper()}; it may be corrupted, malformed, or "
            "password-protected."
        ) from error


def _read_text(path):
    try:
        data = path.read_bytes()
    except OSError as error:
        raise DocumentExtractionError("Unable to read the TXT file.") from error

    try:
        if data.startswith((b"\xff\xfe", b"\xfe\xff")):
            text = data.decode("utf-16")
        else:
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("cp1252")
    except UnicodeDecodeError as error:
        raise DocumentExtractionError(
            "The TXT file uses an unsupported text encoding."
        ) from error

    if "\x00" in text:
        raise DocumentExtractionError(
            "The TXT file contains binary data instead of readable text."
        )
    return text


def _extract_text(path):
    text = _read_text(path)
    return _ensure_content(text, "txt")


def _extract_txt_structured(path):
    text = _read_text(path)
    _ensure_content(text, "txt")
    lines = text.splitlines(keepends=True)
    blocks = []
    current_lines = []
    paragraph_start_line = 0

    def add_paragraph(paragraph_lines, start_line, end_line):
        paragraph_text = "".join(paragraph_lines).rstrip("\r\n")
        if not paragraph_text.strip():
            return
        heading_level = _text_heading_level(paragraph_text)
        blocks.append(_structured_block(
            len(blocks),
            "heading" if heading_level is not None else "paragraph",
            paragraph_text,
            heading_level=heading_level,
            paragraph_index=len(blocks),
            page_start=None,
            page_end=None,
            locator={
                "line_start": start_line + 1,
                "line_end": end_line + 1,
            },
        ))

    def add_heading(line, line_index, heading_level):
        blocks.append(_structured_block(
            len(blocks),
            "heading",
            line.rstrip("\r\n"),
            heading_level=heading_level,
            paragraph_index=len(blocks),
            page_start=None,
            page_end=None,
            locator={
                "line_start": line_index + 1,
                "line_end": line_index + 1,
            },
        ))

    for line_index, line in enumerate(lines):
        line_heading_level = _text_heading_level(line.strip())
        if line_heading_level is not None:
            if current_lines:
                add_paragraph(
                    current_lines,
                    paragraph_start_line,
                    line_index - 1,
                )
                current_lines = []
            add_heading(line, line_index, line_heading_level)
        elif line.strip():
            if not current_lines:
                paragraph_start_line = line_index
            current_lines.append(line)
        elif current_lines:
            add_paragraph(current_lines, paragraph_start_line, line_index - 1)
            current_lines = []
    if current_lines:
        add_paragraph(current_lines, paragraph_start_line, len(lines) - 1)

    return {
        "format": "txt",
        "original_text": text,
        "blocks": blocks,
        "warnings": [],
    }


def extract_structured_document(file_path):
    """Return ordered, source-located content for Docupedia formats only."""
    path = Path(file_path)
    if not path.is_file():
        raise DocumentExtractionError("The uploaded document was not found.")

    file_format = path.suffix.lower()
    if file_format not in DOCPEDIA_FORMATS:
        supported = ", ".join(sorted(DOCPEDIA_FORMATS))
        raise DocumentExtractionError(
            f"Unsupported Docupedia document format '{file_format or '(no extension)'}'. "
            f"Supported formats: {supported}."
        )

    if file_format == ".pdf":
        result = _extract_pdf_structured(path)
    elif file_format == ".docx":
        result = _extract_docx_structured(path)
    elif file_format == ".doc":
        result = _convert_legacy_structured(path, file_format)
    else:
        result = _extract_txt_structured(path)

    if not any(block["block_type"] == "heading" for block in result["blocks"]):
        result["warnings"].append({
            "code": "no_headings_detected",
            "message": (
                "No reliable heading hierarchy was detected; source content "
                "remains in original reading order."
            ),
        })
    result["warnings"] = _deduplicate_warnings(result["warnings"])
    return result


def _deduplicate_warnings(warnings):
    seen = set()
    unique = []
    for warning in warnings:
        key = (
            warning.get("code"),
            warning.get("page_number"),
            warning.get("paragraph_index"),
            warning.get("message"),
        )
        if key not in seen:
            seen.add(key)
            unique.append(warning)
    return unique


def extract_document(file_path):
    """Return readable text from a supported document path.

    Legacy DOC, XLS, and PPT files are converted to their modern formats in
    a temporary directory when LibreOffice is available.
    """
    path = Path(file_path)
    if not path.is_file():
        raise DocumentExtractionError("The uploaded document was not found.")

    file_format = path.suffix.lower()
    if file_format not in SUPPORTED_FORMATS:
        supported = ", ".join(sorted(SUPPORTED_FORMATS))
        raise DocumentExtractionError(
            f"Unsupported document format '{file_format or '(no extension)'}'. "
            f"Supported formats: {supported}."
        )

    if file_format == ".pdf":
        return _extract_pdf(path)
    if file_format == ".docx":
        return _extract_docx(path)
    if file_format in {".doc", ".xls", ".ppt"}:
        return _convert_legacy(path, file_format)
    if file_format == ".xlsx":
        return _extract_xlsx(path)
    if file_format == ".pptx":
        return _extract_pptx(path)
    return _extract_text(path)
